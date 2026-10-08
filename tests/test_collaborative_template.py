import unittest
from pathlib import Path

from docx import Document


TAGS = [
    '{{ descripcion }}',
    '{{ palabras_clave }}',
    '{{r objetivos }}',
    '{{r conceptos_clave }}',
    '{{ resumen }}',
    '{{ metodologia }}',
    '{{ conclusiones }}',
    '{{ discusiones }}',
    '{{r bibliografia }}',
]


class TestCollaborativeTemplate(unittest.TestCase):
    def test_template_contains_expected_tags(self):
        path = Path(__file__).resolve().parents[1] / 'plantilla_protocolo_colaborativo.docx'
        self.assertTrue(path.exists(), 'La plantilla colaborativa no existe')

        document = Document(str(path))
        texts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    texts.append(cell.text)

        combined = '\n'.join(texts)

        for tag in TAGS:
            self.assertIn(tag, combined, f'Falta el tag {tag} en la plantilla colaborativa')

    def test_prompt_builds_without_formatting_error(self):
        from ai_processor import _build_system_prompt
        prompt = _build_system_prompt('colaborativo')
        self.assertIn('"descripcion":', prompt)
        self.assertIn('Protocolo Colaborativo', prompt)


if __name__ == '__main__':
    unittest.main()
